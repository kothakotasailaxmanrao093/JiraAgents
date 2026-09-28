"""Consolidated review-report building: metrics rollup, DOCX rendering, email HTML.

Pure (no network — python-docx only touches an in-memory buffer); kept separate
from the ``build_review_report`` tool so this logic is unit-testable without the
SDK/storage. Mirrors ``review/adf.py``'s split between rendering and the I/O shim
that uploads/posts the result.
"""

from __future__ import annotations

import html
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from .models import FINDING_TYPE_ORDER, READINESS_RANK, ReadinessLevel

_READINESS_COLOR = {
    "ready": RGBColor(0x00, 0x70, 0x27),
    "needs_minor_clarification": RGBColor(0xFF, 0x8C, 0x00),
    "needs_major_clarification": RGBColor(0xCC, 0x33, 0x00),
    "not_ready": RGBColor(0xB0, 0x00, 0x00),
}
READINESS_LABELS = {
    "ready": "Ready",
    "needs_minor_clarification": "Needs Minor Clarification",
    "needs_major_clarification": "Needs Major Clarification",
    "not_ready": "Not Ready",
}


def aggregate_report_metrics(issue_results: list[dict[str, Any]]) -> dict[str, Any]:
    """Roll up per-issue results into report-level metrics: total findings by
    category, error count, and the worst readiness level across all issues."""
    by_category: dict[str, int] = {}
    total_findings = 0
    error_count = 0
    worst_level: ReadinessLevel | None = None

    for r in issue_results:
        for category, count in (r.get("by_category") or {}).items():
            by_category[category] = by_category.get(category, 0) + count
            total_findings += count
        if r.get("status") != "success":
            error_count += 1
        readiness = r.get("readiness")
        if readiness and readiness.get("level"):
            try:
                level = ReadinessLevel(readiness["level"])
            except ValueError:
                level = None
            if level is not None and (
                worst_level is None or READINESS_RANK[level] < READINESS_RANK[worst_level]
            ):
                worst_level = level

    return {
        "issue_count": len(issue_results),
        "error_count": error_count,
        "total_findings": total_findings,
        "by_category": by_category,
        "overall_readiness": worst_level.value if worst_level else "unknown",
    }


def _add_bullets(doc: Document, items: list[str]) -> None:
    if not items:
        doc.add_paragraph("None identified.", style="Intense Quote")
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _finding_line(f: dict[str, Any]) -> str:
    prefix = f"[{f['priority'].upper()}] " if f.get("priority") else ""
    tail_bits = [f"confidence: {f.get('confidence', 'medium')}"]
    if f.get("evidence_source"):
        tail_bits.append(f"evidence: {f['evidence_source']}")
    return f"{prefix}{f.get('description', '')} ({'; '.join(tail_bits)})"


def build_docx_report(
    label: str,
    issue_results: list[dict[str, Any]],
    metrics: dict[str, Any],
    generated_at: str,
) -> bytes:
    """Render the consolidated review report to DOCX bytes."""
    import io

    doc = Document()

    title = doc.add_heading("Jira Requirement Review Report", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    meta = doc.add_paragraph()
    meta.add_run("Issue / Query: ").bold = True
    meta.add_run(f"{label}\n")
    meta.add_run("Generated: ").bold = True
    meta.add_run(f"{generated_at}\n")
    meta.add_run("Issues reviewed: ").bold = True
    meta.add_run(str(metrics.get("issue_count", len(issue_results))))

    doc.add_heading("At a Glance", level=1)
    readiness = metrics.get("overall_readiness", "unknown")
    readiness_para = doc.add_paragraph()
    readiness_para.add_run("Worst Readiness: ").bold = True
    r = readiness_para.add_run(READINESS_LABELS.get(readiness, readiness.title()))
    r.bold = True
    r.font.color.rgb = _READINESS_COLOR.get(readiness, RGBColor(0, 0, 0))

    glance = [f"Total findings: {metrics.get('total_findings', 0)}"]
    for ft in FINDING_TYPE_ORDER:
        count = (metrics.get("by_category") or {}).get(ft.value, 0)
        if count:
            glance.append(f"{ft.value}: {count}")
    if metrics.get("error_count"):
        glance.append(f"Issues with errors/warnings: {metrics['error_count']}")
    _add_bullets(doc, glance)

    for r in issue_results:
        key = r.get("issue_key", "Unknown")
        summary = r.get("summary") or key
        doc.add_heading(f"{key} — {summary}", level=1)

        if r.get("status") != "success" and (r.get("error") or r.get("analysis_error")):
            warn = doc.add_paragraph()
            run = warn.add_run(f"⚠️ {r.get('message') or r.get('error') or r.get('analysis_error')}")
            run.italic = True
            run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
            continue

        issue_readiness = r.get("readiness")
        if issue_readiness and issue_readiness.get("level"):
            level = issue_readiness["level"]
            badge = doc.add_paragraph()
            badge.add_run("Readiness: ").bold = True
            score = issue_readiness.get("score", "?")
            rb = badge.add_run(f"{READINESS_LABELS.get(level, level)} (score {score}/5)")
            rb.bold = True
            rb.font.color.rgb = _READINESS_COLOR.get(level, RGBColor(0, 0, 0))
            if issue_readiness.get("executive_summary"):
                doc.add_paragraph(issue_readiness["executive_summary"])

        findings = r.get("findings") or []
        by_type: dict[str, list[dict[str, Any]]] = {}
        for f in findings:
            by_type.setdefault(f.get("finding_type", ""), []).append(f)

        for ft in FINDING_TYPE_ORDER:
            items = by_type.get(ft.value)
            if not items:
                continue
            doc.add_heading(ft.value, level=2)
            _add_bullets(doc, [_finding_line(f) for f in items])

        if not findings:
            doc.add_paragraph("No findings identified.", style="Intense Quote")

    try:
        doc.styles["Normal"].font.size = Pt(10.5)
    except Exception:
        pass

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _esc(text: Any) -> str:
    return html.escape(str(text if text is not None else ""))


def build_report_email_html(
    label: str, issue_results: list[dict[str, Any]], metrics: dict[str, Any]
) -> str:
    readiness = READINESS_LABELS.get(
        metrics.get("overall_readiness", "unknown"), metrics.get("overall_readiness", "unknown")
    )
    parts: list[str] = [
        f"<h2>Jira Requirement Review Report — {_esc(label)}</h2>",
        "<p>"
        f"<strong>Issues reviewed:</strong> {metrics.get('issue_count', 0)} &nbsp;|&nbsp; "
        f"<strong>Worst readiness:</strong> {_esc(readiness)} &nbsp;|&nbsp; "
        f"<strong>Total findings:</strong> {metrics.get('total_findings', 0)}"
        "</p>",
    ]

    for r in issue_results:
        key = r.get("issue_key", "")
        issue_readiness = r.get("readiness") or {}
        summary = issue_readiness.get("executive_summary") or r.get("message", "")
        parts.append(f"<h3>{_esc(key)} — {_esc(summary)}</h3>")

    parts.append(
        "<p style='color:#666;font-size:12px'>"
        "Full report (all findings, evidence, and confidence/priority tags) available "
        "as a Word document via the download link."
        "</p>"
    )
    return "\n".join(parts)
