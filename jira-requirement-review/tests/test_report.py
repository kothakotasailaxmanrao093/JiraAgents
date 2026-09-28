"""Consolidated review report: metrics rollup, DOCX rendering, email HTML."""

from __future__ import annotations

import io

from docx import Document

from review.report import aggregate_report_metrics, build_docx_report, build_report_email_html


def _issue_result(key, *, status="success", by_category=None, readiness=None, findings=None):
    return {
        "issue_key": key,
        "summary": f"Summary {key}",
        "status": status,
        "by_category": by_category or {},
        "readiness": readiness,
        "findings": findings or [],
    }


def test_aggregate_metrics_sums_categories_and_picks_worst_readiness():
    results = [
        _issue_result(
            "ABC-1",
            by_category={"Assumption": 2, "Edge Case": 1},
            readiness={"level": "ready", "score": 5},
        ),
        _issue_result(
            "ABC-2",
            by_category={"Assumption": 1},
            readiness={"level": "not_ready", "score": 1},
        ),
    ]
    metrics = aggregate_report_metrics(results)
    assert metrics["issue_count"] == 2
    assert metrics["total_findings"] == 4
    assert metrics["by_category"] == {"Assumption": 3, "Edge Case": 1}
    assert metrics["overall_readiness"] == "not_ready"
    assert metrics["error_count"] == 0


def test_aggregate_metrics_counts_non_success_as_errors():
    results = [_issue_result("ABC-1", status="error"), _issue_result("ABC-2", status="success")]
    metrics = aggregate_report_metrics(results)
    assert metrics["error_count"] == 1


def test_aggregate_metrics_handles_missing_readiness():
    metrics = aggregate_report_metrics([_issue_result("ABC-1")])
    assert metrics["overall_readiness"] == "unknown"


def test_build_docx_report_is_valid_and_contains_key_content():
    results = [
        _issue_result(
            "ABC-1",
            by_category={"Ambiguity": 1},
            readiness={
                "level": "needs_major_clarification",
                "score": 2,
                "executive_summary": "Gap.",
            },
            findings=[
                {
                    "finding_type": "Ambiguity",
                    "description": "real-time is undefined",
                    "confidence": "high",
                    "priority": "high",
                    "evidence_source": "Target ABC-1 (description)",
                }
            ],
        )
    ]
    metrics = aggregate_report_metrics(results)
    docx_bytes = build_docx_report("ABC-1", results, metrics, "2026-01-01 00:00 UTC")
    assert docx_bytes

    doc = Document(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Jira Requirement Review Report" in text
    assert "ABC-1" in text
    assert "Needs Major Clarification" in text
    assert "[HIGH] real-time is undefined" in text
    assert "Target ABC-1 (description)" in text


def test_build_docx_report_shows_error_for_failed_issue():
    results = [_issue_result("ABC-2", status="error")]
    results[0]["error"] = "boom"
    results[0]["message"] = "Could not gather context for ABC-2: boom"
    metrics = aggregate_report_metrics(results)
    docx_bytes = build_docx_report("ABC-2", results, metrics, "2026-01-01 00:00 UTC")
    doc = Document(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "boom" in text


def test_build_report_email_html_escapes_and_summarizes():
    results = [
        _issue_result(
            "ABC-1",
            readiness={"level": "ready", "score": 5, "executive_summary": "<script>x</script>"},
        )
    ]
    metrics = aggregate_report_metrics(results)
    email_html = build_report_email_html("ABC-1", results, metrics)
    assert "<h2>Jira Requirement Review Report — ABC-1</h2>" in email_html
    assert "&lt;script&gt;" in email_html  # escaped, not executable
    assert "<script>x</script>" not in email_html
